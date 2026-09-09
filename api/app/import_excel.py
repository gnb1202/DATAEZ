"""Read workbook cell values without pandas numeric inference."""
import io
import posixpath
from datetime import date, datetime
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from .import_validation import ImportValidationError, issue


def _xlsx_rows(content):
    with ZipFile(io.BytesIO(content)) as archive:
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet = book.find("m:sheets/m:sheet", ns)
        rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(r.attrib["Target"] for r in relations if r.attrib["Id"] == rid)
        sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        xml = ET.fromstring(archive.read(sheet_path))
        numbers = {cell.attrib["r"]: cell.find("m:v", ns).text for cell in xml.findall(".//m:sheetData/m:row/m:c", ns)
                   if cell.attrib.get("t", "n") == "n" and cell.find("m:v", ns) is not None}
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        for row in workbook.worksheets[0].iter_rows():
            values, numeric = [], set()
            for i, cell in enumerate(row):
                if cell.data_type in {"f", "e"}:
                    raise ImportValidationError([issue(cell.row, cell.column_letter, "수식·오류 셀은 계산된 값으로 바꿔서 업로드해주세요.")])
                value = cell.value
                if value is not None and cell.data_type == "n" and not isinstance(value, (date, datetime)):
                    # Use the stored XML decimal lexeme, not openpyxl's float.
                    value = numbers.get(cell.coordinate, str(value))
                    numeric.add(i)
                values.append(value)
            yield values, numeric
    finally:
        workbook.close()


def _xls_rows(content):
    try:
        import xlrd
    except ImportError:
        raise ImportValidationError([issue(1, "", "이 서버는 XLS 읽기를 지원하지 않습니다. XLSX 또는 CSV로 저장해주세요.")]) from None
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    for row in range(sheet.nrows):
        values, numeric = [], set()
        for col in range(sheet.ncols):
            cell = sheet.cell(row, col)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, book.datemode)
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                # XLS stores binary doubles; recover their representable value,
                # never claim precision that was absent in the source workbook.
                value = str(value)
                numeric.add(col)
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(value)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                raise ImportValidationError([issue(row + 1, str(col + 1), "오류 셀을 수정해주세요.")])
            values.append(value)
        yield values, numeric


def read_excel(content: bytes, filename: str) -> pd.DataFrame:
    try:
        records = list(_xlsx_rows(content) if filename.lower().endswith(".xlsx") else _xls_rows(content))
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError([issue(1, "", "Excel 파일을 읽지 못했습니다. 파일 형식과 시트를 확인해주세요.")]) from exc
    if not records:
        raise ImportValidationError([issue(1, "", "파일에 데이터가 없습니다.")])
    headers = list(records[0][0])
    while headers and (headers[-1] is None or not str(headers[-1]).strip()):
        headers.pop()
    if not headers or any(h is None or not str(h).strip() for h in headers):
        raise ImportValidationError([issue(1, "", "모든 컬럼에 이름이 필요합니다.")])
    rows, locations, numeric_cells = [], [], set()
    for line, (values, numeric) in enumerate(records[1:], 2):
        if all(v is None or v == "" for v in values):
            continue
        if any(v is not None and v != "" for v in values[len(headers):]):
            raise ImportValidationError([issue(line, "", "이름이 없는 컬럼에 데이터가 있습니다.")])
        numeric_cells.update((len(rows), c) for c in numeric if c < len(headers))
        rows.append(values[:len(headers)])
        locations.append(line)
    df = pd.DataFrame(rows, columns=[str(h) for h in headers], dtype=object)
    df.attrs.update(row_numbers=locations, numeric_cells=numeric_cells)
    return df
